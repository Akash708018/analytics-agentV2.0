"""
Path in, draft spec out.

`preview.py` answers "given these rows, where is the header". This answers
"given this path, what is the whole proposal" -- which sheet, which merges,
which tail, rendered into something a person can read and correct.

It exists so `server.py` keeps no logic. A tool should call one function and
return what it gets back.

Nothing here loads data. The draft is returned as JSON alongside its prose, and
`confirm_ingest_spec` takes that JSON back. That is deliberate: it means the
spec the user confirms is the spec that runs, that they can edit any field
before confirming, and that no draft is held in server memory between calls.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from backend.engine.config import CSV_PREVIEW_LINES, EXCEL_PREVIEW_ROWS
from backend.engine.ingest import csv_loader, excel, merges, preview, sizegate
from backend.engine.ingest.csv_loader import LoadRefused
from backend.engine.ingest.spec import IngestSpec

# Rows read from the end of a sheet when looking for a footer.
TAIL_ROWS = preview.FOOTER_SCAN_ROWS


@dataclass
class Draft:
    spec: IngestSpec | None
    guess: preview.HeaderGuess
    pivot: preview.PivotVerdict
    source_type: str
    sheet: str | None
    sheets: list[str]

    @property
    def needs_answer(self) -> bool:
        return self.spec is None or not self.spec.is_confirmable


def _default_dataset_name(path: Path) -> str:
    stem = "".join(c if c.isalnum() else "_" for c in path.stem).strip("_")
    stem = stem.lower() or "dataset"
    if stem[0].isdigit():
        stem = f"d_{stem}"
    return stem[:63]


def _excel_tail(path: Path, sheet: str) -> list:
    """The last TAIL_ROWS rows of a sheet, for footer detection."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        last = ws.max_row or 0
        if not last:
            return []
        start = max(1, last - TAIL_ROWS + 1)
        return [list(r) for r in ws.iter_rows(min_row=start, values_only=True)]
    finally:
        wb.close()


def _has_values(rows) -> bool:
    return any(v is not None and str(v).strip() for row in rows for v in row)


def _formulas_without_values(path: Path, sheet: str, spec: IngestSpec) -> list[str]:
    """Target names of columns whose data rows (in the preview) hold a formula with no saved value.

    Values are read with data_only=True, so a formula reads as what Excel saved for it. A file
    written by a program (openpyxl, pandas) saves none, and the cell reads None: said here rather
    than loaded as an empty column in silence (P14-O5, B3).
    """
    from openpyxl import load_workbook

    first = spec.data_start_row
    last = first + EXCEL_PREVIEW_ROWS
    found: set[int] = set()
    formulas = load_workbook(path, read_only=True)
    values = load_workbook(path, read_only=True, data_only=True)
    try:
        f_rows = formulas[sheet].iter_rows(min_row=first, max_row=last, values_only=True)
        v_rows = values[sheet].iter_rows(min_row=first, max_row=last, values_only=True)
        for f_row, v_row in zip(f_rows, v_rows):
            for i, (f, v) in enumerate(zip(f_row, v_row)):
                if isinstance(f, str) and f.startswith("=") and v is None:
                    found.add(i)
    finally:
        formulas.close()
        values.close()
    names = [c.target_name for c in spec.columns]
    return [names[i] for i in sorted(found) if i < len(names)]


def draft_for_path(
    path: str | Path,
    dataset_name: str | None = None,
    sheet: str | None = None,
    header_rows: list[int] | None = None,
    header_join: str | None = None,
    authorised_fill: bool = False,
) -> Draft:
    """
    Build a draft spec for one file. Reads only the top and tail.

    `header_rows`, `header_join` and `authorised_fill` carry a person's answer
    back in. Supplying `header_rows` settles an ambiguous file, which is the
    only way `unresolved` gets cleared.
    """
    p = Path(path)
    if not p.exists():
        raise LoadRefused(
            f"BLOCKED: no file at {path}.\n"
            f"NEXT STEP: check the path. Relative paths resolve against the "
            f"server's working directory, not your shell's, so an absolute "
            f"path is safer."
        )

    gate = sizegate.check_file(p)
    if not gate.allowed:
        raise LoadRefused(gate.message)

    name = dataset_name or _default_dataset_name(p)
    source_type = sizegate.source_type_for(p)

    if source_type == "excel":
        try:
            sheets = excel.list_sheets(p)
        except Exception as exc:  # noqa: BLE001 - BadZipFile, InvalidFileException, KeyError
            # A CSV renamed .xlsx, a truncated download, a password-protected workbook: all
            # arrive here as a file that is not a zip of sheets (P14-D70).
            raise LoadRefused(
                f"BLOCKED: {p.name} is not a readable .xlsx workbook.\n"
                f"WHY: {type(exc).__name__}: {exc}. It may be another kind of file renamed, "
                f"cut short, or protected with a password.\n"
                f"NEXT STEP: open it in Excel and save it as .xlsx, or export it as CSV."
            ) from exc
        chosen = sheet or merges.active_sheet_name(str(p))
        if chosen not in sheets:
            raise LoadRefused(
                f"BLOCKED: {p.name} has no sheet called {chosen!r}.\n"
                f"Sheets present: {', '.join(sheets)}\n"
                f"NEXT STEP: pass one of those names."
            )
        rows = excel.preview_rows(p, chosen, n=EXCEL_PREVIEW_ROWS)
        skipped_note = None
        if not _has_values(rows):
            # An empty sheet, as a cover page often is. Named explicitly, that is a refusal that
            # lists the sheets holding data; not named, the first sheet with data is read and the
            # draft says so. It used to raise "No preview rows" (P14-O7, B5).
            with_data = [s for s in sheets
                         if s != chosen and _has_values(excel.preview_rows(p, s, n=5))]
            if sheet is not None or not with_data:
                raise LoadRefused(
                    f"BLOCKED: sheet {chosen!r} of {p.name} is empty.\n"
                    f"WHY: there is no row in it to read a header or data from.\n"
                    + (f"Sheets holding data: {', '.join(with_data)}\n"
                       f"NEXT STEP: pass sheet={with_data[0]!r}." if with_data else
                       "NEXT STEP: check the workbook -- no sheet in it holds data.")
                )
            skipped_note = (f"Sheet {chosen!r} is empty, so {with_data[0]!r} -- the first sheet "
                            f"holding data -- is read instead. Name another sheet to change it.")
            chosen = with_data[0]
            rows = excel.preview_rows(p, chosen, n=EXCEL_PREVIEW_ROWS)
        spec, guess, pivot = preview.draft_spec(
            rows,
            path=str(p),
            source_type="excel",
            dataset_name=name,
            sheet=chosen,
            merge_refs=merges.merged_ranges(str(p), chosen),
            tail_rows=_excel_tail(p, chosen),
            header_rows=header_rows,
            header_join=header_join,
        )
        if spec is not None:
            if skipped_note:
                spec.assumptions.insert(0, skipped_note)
            unsaved = _formulas_without_values(p, chosen, spec)
            if unsaved:
                spec.assumptions.append(
                    f"{', '.join(unsaved)} hold{'s' if len(unsaved) == 1 else ''} formulas with "
                    f"no saved result: the file was written by a program, not saved by Excel, so "
                    f"there is no value to read and those cells load empty. Open the file in "
                    f"Excel, save it, and upload it again to load the computed values.")
        return Draft(spec, guess, pivot, "excel", chosen, sheets)

    lines = csv_loader.preview_lines(p, n=CSV_PREVIEW_LINES)
    # Lines starting with '#' above everything else are comments ('# exported by ...'): the
    # header is the first line after them, not a question to ask (P14-D54).
    comments = next((i for i, ln in enumerate(lines) if not ln.lstrip().startswith("#")),
                    len(lines))
    comment_note = None
    if comments and header_rows is None and comments < len(lines):
        header_rows = [comments + 1]
        comment_note = (f"The first {comments} line(s) begin with '#', so they are comments; "
                        f"line {comments + 1} is the header.")
    rows = preview.parse_csv_preview(lines[comments:] if comments else lines)
    rows = [[] for _ in range(comments)] + rows
    # The tail is parsed with the head's delimiter: twenty lines of totals and notes are too
    # few to sniff one from.
    delimiter = preview.sniff_delimiter(lines)
    tail = [r for r in preview.parse_csv_preview(csv_loader.tail_lines(p, TAIL_ROWS),
                                                 delimiter=delimiter)
            if any(str(v).strip() for v in r)][-TAIL_ROWS:]
    spec, guess, pivot = preview.draft_spec(
        rows, path=str(p), source_type="csv", dataset_name=name,
        header_rows=header_rows, header_join=header_join,
        authorised_fill=authorised_fill, tail_rows=tail,
    )
    if spec is not None and comment_note:
        spec.assumptions = [comment_note] + [
            a for a in spec.assumptions if not a.startswith("header_rows was given")]
    return Draft(spec, guess, pivot, "csv", None, [])


def render(draft: Draft) -> str:
    """
    The draft as the user reads it.

    Ordered so the questions come last: whatever is at the bottom of a tool
    result is what gets answered.
    """
    out: list[str] = []

    if draft.source_type == "excel" and len(draft.sheets) > 1:
        others = [s for s in draft.sheets if s != draft.sheet]
        out.append(
            f"Reading sheet '{draft.sheet}'. Also in this workbook: "
            f"{', '.join(others)}."
        )
        out.append("")

    if draft.spec is None:
        out.append("No spec proposed -- this file is ambiguous.")
        out.append("")
        out.append("What was worked out:")
        out += [f"  - {r}" for r in draft.guess.reasons]
        out.append("")
        out.append("What has to be settled first:")
        out += [f"  - {q}" for q in draft.guess.questions]
        out.append("")
        out.append(
            "Answer and I will build the spec. Nothing has been loaded and "
            "nothing has been assumed."
        )
        return "\n".join(out)

    # to_text already prints the assumptions. Printing them again here read
    # as two different lists on the first run, which is worse than terse.
    out.append(draft.spec.to_text())
    out.append("")
    out.append(f"Confidence: {draft.guess.confidence}")
    if draft.spec.assumptions:
        out.append("Every assumption above is yours to overrule.")

    if draft.spec.unresolved:
        out.append("")
        out.append("THIS SPEC CANNOT BE LOADED AS IT STANDS.")
        out.append(
            f"{', '.join(draft.spec.unresolved)} below is a default, not a "
            f"reading of the file. What is shown is what a load WOULD do."
        )
        out.append("")
        out.append("Put these to the user:")
        out += [f"  - {q}" for q in draft.spec.questions]
        out.append("")
        out.append(
            "Then call propose_ingest_spec again with their answer -- "
            "header_rows=[...], and header_join or authorised_fill if "
            "relevant. It returns a spec with nothing outstanding, and that "
            "one can be confirmed."
        )
        return "\n".join(out)

    # spec.questions, NOT guess.questions. The guess describes what the FILE
    # could settle and never changes; the spec describes what is still
    # outstanding after a person has spoken. Reading the guess here printed
    # "Open questions" at someone who had just answered them.
    if draft.spec.questions:
        out.append("")
        out.append("Open questions:")
        out += [f"  - {q}" for q in draft.spec.questions]

    out.append("")
    out.append(
        "Nothing is loaded yet. To load exactly this, pass the JSON below back "
        "to confirm_ingest_spec. To change something, edit a field first -- "
        "the spec that comes back is the spec that runs."
    )
    out.append("")
    out.append("```json")
    out.append(draft.spec.model_dump_json(indent=2))
    out.append("```")
    return "\n".join(out)


def spec_from_json(spec_json: str) -> IngestSpec:
    """Parse a confirmed spec, turning any validation error into a refusal."""
    text = spec_json.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[1] if "\n" in text else ""
        text = text.rsplit("```", 1)[0]
    try:
        return IngestSpec.model_validate_json(text)
    except Exception as exc:
        raise LoadRefused(
            f"BLOCKED: that is not a usable ingest spec.\n"
            f"Detail: {exc}\n"
            f"NEXT STEP: call propose_ingest_spec to get a valid one, edit the "
            f"fields you want to change, and pass the whole JSON object back."
        ) from exc


__all__ = ["Draft", "draft_for_path", "render", "spec_from_json"]
