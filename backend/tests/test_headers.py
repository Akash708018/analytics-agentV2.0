"""
Unit tests for ingest/headers.py.

Fixtures are built in tmp_path so the suite is self-contained. The one
exception in spirit is `sales_shaped`, which deliberately mirrors the shape of
tests/fixtures/merged_multiheader.xlsx -- sheet 'Sales', 8 columns A-H, header
block rows 1-2, merges A1:B1, C1:E1, F1:H1 -- so a change to that fixture shows
up here rather than three steps later.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from backend.engine.ingest.headers import (
    assemble_names,
    to_target_name,
    to_target_names,
)
from backend.engine.ingest.merges import merged_ranges


@pytest.fixture
def sales_shaped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Order"
    ws["C1"] = "Q3 2024"
    ws["F1"] = "Q4 2024"
    leaves = ["ID", "Region", "Units", "Amt", "Channel", "Units", "Amt", "Channel"]
    for col, val in zip("ABCDEFGH", leaves):
        ws[f"{col}2"] = val
    ws.merge_cells("A1:B1")
    ws.merge_cells("C1:E1")
    ws.merge_cells("F1:H1")
    ws.append(["A1", "North", 10, 100, "web", 11, 110, "store"])
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    return str(path)


def _read_header(path, sheet, last_row):
    wb = load_workbook(path, read_only=True)
    rows = [
        list(r)
        for r in wb[sheet].iter_rows(min_row=1, max_row=last_row, values_only=True)
    ]
    wb.close()
    return rows


# --------------------------------------------------------------------------
# assembly on a real workbook
# --------------------------------------------------------------------------

def test_two_row_excel_header_assembles_end_to_end(sales_shaped):
    refs = merged_ranges(sales_shaped, "Sales")
    rows = _read_header(sales_shaped, "Sales", 2)

    r = assemble_names(
        rows, header_rows_1idx=[1, 2], source_type="excel", merge_refs=refs
    )

    assert r.names == [
        "Order ID",
        "Order Region",
        "Q3 2024 Units",
        "Q3 2024 Amt",
        "Q3 2024 Channel",
        "Q4 2024 Units",
        "Q4 2024 Amt",
        "Q4 2024 Channel",
    ]
    assert r.column_count == 8
    assert r.filled_merges == ["A1:B1", "C1:E1", "F1:H1"]
    assert r.renamed_duplicates == []
    assert r.blank_columns == []


def test_spanning_label_does_not_cross_into_the_next_span(sales_shaped):
    """Q3 must not appear on a Q4 column, and vice versa."""
    refs = merged_ranges(sales_shaped, "Sales")
    rows = _read_header(sales_shaped, "Sales", 2)
    names = assemble_names(
        rows, header_rows_1idx=[1, 2], source_type="excel", merge_refs=refs
    ).names

    assert not any("Q3" in n and "Q4" in n for n in names)
    assert [n for n in names if n.startswith("Q3")] == [
        "Q3 2024 Units",
        "Q3 2024 Amt",
        "Q3 2024 Channel",
    ]


def test_note_counts_blanks_before_the_fill(sales_shaped):
    """
    The number that explains why filling was needed is the count in the row as
    read, not the count after filling -- which is always zero.
    """
    refs = merged_ranges(sales_shaped, "Sales")
    rows = _read_header(sales_shaped, "Sales", 2)
    note = assemble_names(
        rows, header_rows_1idx=[1, 2], source_type="excel", merge_refs=refs
    ).notes[0]

    assert "5 of 8 columns blank" in note
    assert "A1:B1, C1:E1, F1:H1" in note


def test_note_does_not_claim_the_fill_was_used_under_bottom_only(sales_shaped):
    """
    Under bottom_only every row above the last is discarded, so the merged
    labels never reach the names. Saying they were "used to build column
    names" would be false, and an assumptions list is only worth reading if it
    can be trusted without checking.
    """
    refs = merged_ranges(sales_shaped, "Sales")
    rows = _read_header(sales_shaped, "Sales", 2)
    note = assemble_names(
        rows,
        header_rows_1idx=[1, 2],
        source_type="excel",
        join="bottom_only",
        merge_refs=refs,
    ).notes[0]

    assert "filled but not used" in note
    assert "do not appear" in note
    assert "A1:B1, C1:E1, F1:H1" in note


def test_note_still_claims_the_fill_was_used_under_space(sales_shaped):
    refs = merged_ranges(sales_shaped, "Sales")
    rows = _read_header(sales_shaped, "Sales", 2)
    note = assemble_names(
        rows, header_rows_1idx=[1, 2], source_type="excel", join="space",
        merge_refs=refs,
    ).notes[0]

    assert "Filled them to build column names" in note
    assert "filled but not used" not in note


# --------------------------------------------------------------------------
# join modes
# --------------------------------------------------------------------------

ROWS = [["Q3 2024", None, "Q4 2024", None], ["Units", "Amt", "Units", "Amt"]]
MERGES = ["A1:B1", "C1:D1"]


@pytest.mark.parametrize(
    "join,expected",
    [
        ("space", ["Q3 2024 Units", "Q3 2024 Amt", "Q4 2024 Units", "Q4 2024 Amt"]),
        ("underscore", ["Q3 2024_Units", "Q3 2024_Amt", "Q4 2024_Units", "Q4 2024_Amt"]),
        ("bottom_only", ["Units", "Amt", "Units_2", "Amt_2"]),
        ("top_only", ["Q3 2024", "Q3 2024_2", "Q4 2024", "Q4 2024_2"]),
    ],
)
def test_all_four_join_modes(join, expected):
    r = assemble_names(
        [list(x) for x in ROWS],
        header_rows_1idx=[1, 2],
        source_type="excel",
        join=join,
        merge_refs=MERGES,
    )
    assert r.names == expected


def test_unknown_join_mode_lists_the_valid_ones():
    with pytest.raises(ValueError) as exc:
        assemble_names(
            [["a"]], header_rows_1idx=[1], source_type="excel", join="pipe"
        )
    assert "bottom_only" in str(exc.value)


# --------------------------------------------------------------------------
# dedup and blanks
# --------------------------------------------------------------------------

def test_duplicates_get_numeric_suffixes_and_are_reported():
    r = assemble_names(
        [["units", "units", "units"]], header_rows_1idx=[1], source_type="csv"
    )
    assert r.names == ["units", "units_2", "units_3"]
    assert r.renamed_duplicates == [("units", "units_2"), ("units", "units_3")]
    assert "units -> units_2" in r.notes[-1]


def test_dedup_skips_a_suffix_that_is_already_taken():
    """An existing 'units_2' must not be collided with."""
    r = assemble_names(
        [["units", "units_2", "units"]], header_rows_1idx=[1], source_type="csv"
    )
    assert r.names == ["units", "units_2", "units_3"]


def test_blank_columns_are_named_by_position_and_reported():
    r = assemble_names([["a", None, "c"]], header_rows_1idx=[1], source_type="csv")
    assert r.names == ["a", "column_2", "c"]
    assert r.blank_columns == [1]
    assert "column_N" in r.notes[-1]


def test_whitespace_only_cells_count_as_blank():
    r = assemble_names([["a", "   ", "c"]], header_rows_1idx=[1], source_type="csv")
    assert r.names[1] == "column_2"


def test_ragged_header_rows_are_padded_to_the_widest():
    r = assemble_names(
        [["Q3", None], ["Units", "Amt", "Channel"]],
        header_rows_1idx=[1, 2],
        source_type="csv",
    )
    assert r.names == ["Q3 Units", "Amt", "Channel"]


# --------------------------------------------------------------------------
# the CSV rule -- locked decision 10
# --------------------------------------------------------------------------

def test_csv_header_is_never_filled():
    """The same rows an Excel sheet would fill must stay unfilled for a CSV."""
    rows = [["Q3 2024", None, "Q4 2024", None], ["Units", "Amt", "Units", "Amt"]]
    r = assemble_names(rows, header_rows_1idx=[1, 2], source_type="csv")
    assert r.names == ["Q3 2024 Units", "Amt", "Q4 2024 Units", "Amt_2"]
    assert r.filled_merges == []


def test_csv_blanks_in_upper_rows_are_flagged_as_ambiguous():
    rows = [["Q3 2024", None, "Q4 2024", None], ["Units", "Amt", "Units", "Amt"]]
    r = assemble_names(rows, header_rows_1idx=[1, 2], source_type="csv")
    assert r.ambiguous_blanks == [1, 3]
    note = " ".join(r.notes)
    assert "no merge information" in note
    assert "2, 4" in note


def test_single_row_csv_header_raises_no_ambiguity():
    r = assemble_names([["a", "b", "c"]], header_rows_1idx=[1], source_type="csv")
    assert r.ambiguous_blanks == []


def test_passing_merge_refs_for_a_csv_is_refused():
    with pytest.raises(ValueError) as exc:
        assemble_names(
            [["a", None]], header_rows_1idx=[1], source_type="csv", merge_refs=["A1:B1"]
        )
    assert "never auto-filled" in str(exc.value)


# --------------------------------------------------------------------------
# argument validation
# --------------------------------------------------------------------------

def test_row_count_must_match_row_numbers():
    with pytest.raises(ValueError) as exc:
        assemble_names(
            [["a"], ["b"]], header_rows_1idx=[1], source_type="excel"
        )
    assert "must correspond" in str(exc.value)


def test_empty_header_is_refused():
    with pytest.raises(ValueError):
        assemble_names([], header_rows_1idx=[], source_type="excel")


def test_unknown_source_type_is_refused():
    with pytest.raises(ValueError):
        assemble_names([["a"]], header_rows_1idx=[1], source_type="parquet")


# --------------------------------------------------------------------------
# target-name normalisation
# --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "source,expected",
    [
        ("Q3 2024 Amt", "q3_2024_amt"),
        ("Order ID", "order_id"),
        ("  Region  ", "region"),
        ("% of total", "pct_of_total"),
        ("# orders", "num_orders"),
        ("R&D spend", "r_and_d_spend"),
        ("2024", "col_2024"),
        ("amount (INR)", "amount_inr"),
        ("", "unnamed"),
        ("---", "unnamed"),
    ],
)
def test_to_target_name(source, expected):
    assert to_target_name(source) == expected


def test_target_names_dedupe_after_normalising():
    """'Units' and 'units' normalise to the same thing and must not collide."""
    assert to_target_names(["Units", "units", "UNITS"]) == [
        "units",
        "units_2",
        "units_3",
    ]


def test_target_names_survive_the_sales_header():
    names = [
        "Order ID",
        "Order Region",
        "Q3 2024 Units",
        "Q3 2024 Amt",
        "Q3 2024 Channel",
        "Q4 2024 Units",
        "Q4 2024 Amt",
        "Q4 2024 Channel",
    ]
    assert to_target_names(names) == [
        "order_id",
        "order_region",
        "q3_2024_units",
        "q3_2024_amt",
        "q3_2024_channel",
        "q4_2024_units",
        "q4_2024_amt",
        "q4_2024_channel",
    ]
