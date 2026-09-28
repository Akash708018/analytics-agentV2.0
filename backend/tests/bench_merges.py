"""
One-off memory check for merges.merged_ranges.

Not a unit test -- peak-memory assertions are flaky across machines. Run it
by hand once, record the number, move on.
"""

import sys
import time
import tracemalloc
from pathlib import Path

from openpyxl import Workbook

from backend.engine.ingest.merges import merged_ranges

PATH = Path("/tmp/merge_bench.xlsx")
ROWS = 200_000


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Big"
    ws["A1"] = "Region"
    ws["C1"] = "Chan"
    ws.merge_cells("A1:B1")
    for i in range(2, ROWS + 2):
        ws.append([f"r{i % 50}", i, "web"])
    wb.save(PATH)


if __name__ == "__main__":
    if not PATH.exists():
        print(f"building {ROWS:,}-row fixture at {PATH} ...")
        build()
    size_mb = PATH.stat().st_size / 1e6
    tracemalloc.start()
    t0 = time.time()
    refs = merged_ranges(str(PATH), "Big")
    elapsed = time.time() - t0
    _cur, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    print(f"workbook      : {size_mb:.1f} MB on disk, {ROWS:,} rows")
    print(f"merges found  : {refs}")
    print(f"peak memory   : {peak / 1e6:.1f} MB")
    print(f"elapsed       : {elapsed:.2f} s")
    print()
    print("PASS" if peak / 1e6 < 50 else "FAIL -- peak memory is not flat")
    sys.exit(0 if peak / 1e6 < 50 else 1)
